from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parent
SLIDES_DIR = ROOT / "slides"
OUTPUT = ROOT / "final" / "iRent_InspectFlow_重新設計版.pptx"


def main() -> None:
    images = sorted(SLIDES_DIR.glob("slide-*.png"))
    if len(images) != 16:
        raise SystemExit(f"Expected 16 slide images, found {len(images)}")

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for image_path in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

    if len(prs.slides) > len(images):
        slide_id = prs.slides._sldIdLst[0]
        prs.part.drop_rel(slide_id.rId)
        del prs.slides._sldIdLst[0]

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Wrote {len(images)} slides to {OUTPUT}")


if __name__ == "__main__":
    main()
